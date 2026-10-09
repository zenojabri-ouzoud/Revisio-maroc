import streamlit as st
import hashlib
import sqlite3
import os
import base64
import random
import time
import json
import datetime
import pathlib
import re
import html

# =========================================================
# 3AC RevisioMaroc — Application complète
# =========================================================

st.set_page_config(
    page_title="3AC RevisioMaroc",
    page_icon="📚",
    layout="wide",
    initial_sidebar_state="expanded"
)

DB_PATH = "revisiomaroc.db"
UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

SUBJECTS = {
    "maths": {"name": "الرياضيات", "icon": "📐", "color": "#4A90E2"},
    "french": {"name": "اللغة الفرنسية", "icon": "🇫🇷", "color": "#E74C3C"},
    "english": {"name": "اللغة الإنجليزية", "icon": "🇬🇧", "color": "#3498DB"},
    "history": {"name": "الاجتماعيات", "icon": "🌍", "color": "#F39C12"},
    "islamic": {"name": "التربية الإسلامية", "icon": "🕌", "color": "#27AE60"},
    "pc": {"name": "الفيزياء والكيمياء", "icon": "⚗️", "color": "#9B59B6"},
    "svt": {"name": "علوم الحياة والأرض", "icon": "🧬", "color": "#16A085"},
}

THEMES = {
    "⚽ FC Barcelona": {
        "bg": "#0A1E3F", "card": "#1A2F5C", "text": "#F0F8FF",
        "accent": "#A50044", "secondary": "#0F2A52",
        "border": "#004D98", "highlight": "#00D26A"
    },
    "👑 Real Madrid": {
        "bg": "#0F1B2D", "card": "#1E2E4A", "text": "#FFFFFF",
        "accent": "#FEBE10", "secondary": "#0A1421",
        "border": "#00529F", "highlight": "#FEBE10"
    },
    "🦅 الأهلي": {
        "bg": "#1A0A0A", "card": "#2D1515", "text": "#FFF5F5",
        "accent": "#E30613", "secondary": "#120505",
        "border": "#8B0000", "highlight": "#FFD700"
    },
    "🌙 Moonlight": {
        "bg": "#0A0E1A", "card": "#1A1F35", "text": "#E8ECFF",
        "accent": "#7B9FFF", "secondary": "#050810",
        "border": "#3D4A7A", "highlight": "#FFE57F"
    },
    "🌙 Midnight Purple": {
        "bg": "#0D0B1F", "card": "#1A1735", "text": "#EDE9FE",
        "accent": "#A78BFA", "secondary": "#221D4A",
        "border": "#3D3475", "highlight": "#4ADE80"
    },
    "🌊 Ocean Deep": {
        "bg": "#0A1929", "card": "#132F4C", "text": "#E3F2FD",
        "accent": "#00B8D4", "secondary": "#0F2537",
        "border": "#1E4976", "highlight": "#4ADE80"
    },
    "📚 Study Mode": {
        "bg": "#1A1410", "card": "#2D2418", "text": "#FFF8E7",
        "accent": "#D4A574", "secondary": "#0F0B07",
        "border": "#5C4A2E", "highlight": "#FFD700"
    },
    "🌅 Golden Sunset": {
        "bg": "#1F1410", "card": "#331F17", "text": "#FFF3E0",
        "accent": "#FFB74D", "secondary": "#2A1A12",
        "border": "#5D3A24", "highlight": "#4ADE80"
    },
    "🌿 Forest Emerald": {
        "bg": "#0B1F14", "card": "#143728", "text": "#E8F5E9",
        "accent": "#4ADE80", "secondary": "#0F2A1D",
        "border": "#1E5C3D", "highlight": "#00D26A"
    },
    "☀️ Light Mode": {
        "bg": "#F5F7FA", "card": "#FFFFFF", "text": "#1A202C",
        "accent": "#4A90E2", "secondary": "#E2E8F0",
        "border": "#CBD5E0", "highlight": "#38A169"
    },
    "🌌 Galaxy": {
        "bg": "#0D0221", "card": "#1A0533", "text": "#E8D5FF",
        "accent": "#C77DFF", "secondary": "#050011",
        "border": "#7209B7", "highlight": "#4CC9F0"
    },
}

BADGES = {
    "first_quiz": "🎯 أول اختبار",
    "perfect": "💯 العلامة الكاملة",
    "5_quizzes": "🔥 مجتهد",
    "10_quizzes": "💪 مثابر",
    "25_quizzes": "🏃 عدّاء",
    "50_quizzes": "🚀 صاروخ",
    "100_quizzes": "🌟 أسطورة",
    "level_5": "⭐ نجم",
    "level_10": "🌟 نجم لامع",
    "level_20": "👑 ملك",
    "level_50": "🏆 أسطورة حية",
    "all_subjects": "🎓 الموسوعي",
    "streak_7": "🔥 أسبوع كامل",
    "streak_30": "🌋 شهر كامل",
}

RANKS = [
    (50, "🌟 أسطورة حية"),
    (35, "🏆 أسطورة"),
    (20, "👑 خبير"),
    (12, "🏅 متفوق"),
    (8, "⭐ متميز"),
    (5, "🎯 مجتهد"),
    (3, "📖 متعلّم"),
    (1, "🌱 مبتدئ"),
]

TRANSLATIONS = {
    "fr": {
        "dashboard": "Tableau de bord",
        "lessons": "Leçons",
        "quiz": "Quiz",
        "review": "Révision rapide",
        "flashcards": "Cartes mémoire",
        "plan": "Plan d'étude",
        "leaderboard": "Classement",
        "stats": "Mes statistiques",
        "favorites": "Favoris",
        "notifications": "Notifications",
        "developer": "Espace développeur",
        "logout": "Déconnexion",
        "language": "Langue",
        "theme": "Thème",
    },
    "en": {
        "dashboard": "Dashboard",
        "lessons": "Lessons",
        "quiz": "Quiz",
        "review": "Quick revision",
        "flashcards": "Flashcards",
        "plan": "Study plan",
        "leaderboard": "Leaderboard",
        "stats": "My statistics",
        "favorites": "Favorites",
        "notifications": "Notifications",
        "developer": "Developer panel",
        "logout": "Log out",
        "language": "Language",
        "theme": "Theme",
    }
}


def T(key):
    lang = st.session_state.get("language", "ar")
    if lang == "ar":
        arabic = {
            "dashboard": "الرئيسية",
            "lessons": "الدروس",
            "quiz": "الاختبارات",
            "review": "المراجعة السريعة",
            "flashcards": "بطاقات الحفظ",
            "plan": "خطة الدراسة",
            "leaderboard": "لوحة المتصدرين",
            "stats": "إحصائياتي",
            "favorites": "المفضلة",
            "notifications": "الإشعارات",
            "developer": "فضاء المطور",
            "logout": "تسجيل الخروج",
            "language": "اللغة",
            "theme": "الثيم",
        }
        return arabic.get(key, key)
    return TRANSLATIONS.get(lang, {}).get(key, key)


# =========================================================
# قاعدة البيانات
# =========================================================

def get_db():
    conn = sqlite3.connect(DB_PATH, timeout=20)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def hash_password(p):
    return hashlib.sha256(str(p).encode("utf-8")).hexdigest()


def init_db():
    conn = get_db()
    cur = conn.cursor()

    statements = [
        """CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            role TEXT DEFAULT 'student',
            full_name TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )""",
        """CREATE TABLE IF NOT EXISTS lessons (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            subject TEXT NOT NULL,
            language TEXT DEFAULT 'ar',
            title TEXT NOT NULL,
            content TEXT DEFAULT '',
            image_url TEXT DEFAULT '',
            pdf_url TEXT DEFAULT '',
            owner TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )""",
        """CREATE TABLE IF NOT EXISTS questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            lesson_id INTEGER NOT NULL,
            question TEXT NOT NULL,
            option_a TEXT NOT NULL,
            option_b TEXT NOT NULL,
            option_c TEXT NOT NULL,
            option_d TEXT NOT NULL,
            correct_answer TEXT NOT NULL,
            explanation TEXT DEFAULT '',
            FOREIGN KEY(lesson_id) REFERENCES lessons(id) ON DELETE CASCADE
        )""",
        """CREATE TABLE IF NOT EXISTS quiz_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            lesson_id INTEGER,
            lesson_title TEXT,
            subject TEXT,
            score INTEGER,
            total INTEGER,
            percent REAL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )""",
        """CREATE TABLE IF NOT EXISTS favorites (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            lesson_id INTEGER NOT NULL,
            UNIQUE(username, lesson_id)
        )""",
        """CREATE TABLE IF NOT EXISTS user_stats (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            total_points INTEGER DEFAULT 0,
            level INTEGER DEFAULT 1,
            quizzes_taken INTEGER DEFAULT 0,
            perfect_scores INTEGER DEFAULT 0,
            unique_subjects TEXT DEFAULT '[]',
            badges TEXT DEFAULT '[]',
            last_daily TEXT DEFAULT '',
            streak INTEGER DEFAULT 0
        )""",
        """CREATE TABLE IF NOT EXISTS reviews (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            lesson_id INTEGER NOT NULL,
            username TEXT NOT NULL,
            rating INTEGER NOT NULL,
            comment TEXT DEFAULT '',
            UNIQUE(lesson_id, username)
        )""",
        """CREATE TABLE IF NOT EXISTS messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            lesson_id INTEGER NOT NULL,
            username TEXT NOT NULL,
            message TEXT NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )""",
        """CREATE TABLE IF NOT EXISTS notifications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            title TEXT,
            message TEXT,
            icon TEXT DEFAULT '🔔',
            is_read INTEGER DEFAULT 0
        )""",
        """CREATE TABLE IF NOT EXISTS flashcards (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            subject TEXT,
            front TEXT,
            back TEXT,
            known INTEGER DEFAULT 0
        )""",
        """CREATE TABLE IF NOT EXISTS notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            lesson_id INTEGER NOT NULL,
            note TEXT NOT NULL
        )""",
        """CREATE TABLE IF NOT EXISTS friends (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            friend_username TEXT NOT NULL,
            status TEXT DEFAULT 'pending',
            UNIQUE(username, friend_username)
        )""",
        """CREATE TABLE IF NOT EXISTS study_plan (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            subject TEXT NOT NULL,
            priority TEXT DEFAULT 'متوسطة',
            target_date TEXT,
            completed INTEGER DEFAULT 0
        )"""
    ]

    for statement in statements:
        cur.execute(statement)

    # ترقية بسيطة للجداول التي أنشأتها إصدارات سابقة من التطبيق.
    # تضمن الأعمدة الاختيارية المطلوبة عند تحديث قاعدة قديمة.
    upgrades = {
        "lessons": {
            "language": "TEXT DEFAULT 'ar'",
            "image_url": "TEXT DEFAULT ''",
            "pdf_url": "TEXT DEFAULT ''",
            "owner": "TEXT",
            "created_at": "TEXT"
        },
        "quiz_history": {
            "lesson_title": "TEXT",
            "subject": "TEXT",
            "percent": "REAL",
            "created_at": "TEXT"
        },
        "user_stats": {
            "unique_subjects": "TEXT DEFAULT '[]'",
            "badges": "TEXT DEFAULT '[]'",
            "last_daily": "TEXT DEFAULT ''",
            "streak": "INTEGER DEFAULT 0"
        },
        "messages": {"created_at": "TEXT"}
    }

    for table, columns in upgrades.items():
        existing = {
            row["name"] for row in
            cur.execute(f"PRAGMA table_info({table})").fetchall()
        }
        for column, definition in columns.items():
            if column not in existing:
                cur.execute(
                    f"ALTER TABLE {table} ADD COLUMN {column} {definition}"
                )

    cur.execute(
        """INSERT OR IGNORE INTO users
        (username, password_hash, role, full_name)
        VALUES (?, ?, ?, ?)""",
        ("soufianeDEV", hash_password("soufiane2030"),
         "developer", "Soufiane Ouhazza")
    )

    cur.execute(
        """INSERT OR IGNORE INTO user_stats
        (username, total_points, level)
        VALUES (?, 0, 1)""",
        ("soufianeDEV",)
    )

    conn.commit()
    conn.close()


def register_user(u, p, n):
    u = u.strip()
    n = n.strip()
    if len(u) < 3 or len(p) < 6 or not n:
        return False, "اسم المستخدم يجب أن يحتوي على 3 أحرف على الأقل، وكلمة المرور على 6 أحرف."
    if u.lower() == "soufianedev":
        return False, "هذا الاسم محجوز."
    conn = get_db()
    try:
        conn.execute(
            """INSERT INTO users
            (username, password_hash, role, full_name)
            VALUES (?, ?, 'student', ?)""",
            (u, hash_password(p), n)
        )
        conn.execute(
            "INSERT INTO user_stats (username) VALUES (?)", (u,)
        )
        conn.commit()
        return True, "تم إنشاء الحساب بنجاح."
    except sqlite3.IntegrityError:
        return False, "اسم المستخدم مستعمل من قبل."
    finally:
        conn.close()


def authenticate(u, p):
    conn = get_db()
    row = conn.execute(
        "SELECT * FROM users WHERE username=? AND password_hash=?",
        (u.strip(), hash_password(p))
    ).fetchone()
    conn.close()
    return dict(row) if row else None


def get_user_stats(u):
    conn = get_db()
    row = conn.execute(
        "SELECT * FROM user_stats WHERE username=?", (u,)
    ).fetchone()
    conn.close()
    if row:
        return dict(row)
    conn = get_db()
    conn.execute(
        "INSERT OR IGNORE INTO user_stats (username) VALUES (?)", (u,)
    )
    conn.commit()
    row = conn.execute(
        "SELECT * FROM user_stats WHERE username=?", (u,)
    ).fetchone()
    conn.close()
    return dict(row)


def update_user_stats(u, points, quiz, perfect, subject):
    conn = get_db()
    conn.execute(
        "INSERT OR IGNORE INTO user_stats (username) VALUES (?)", (u,)
    )
    row = conn.execute(
        "SELECT * FROM user_stats WHERE username=?", (u,)
    ).fetchone()
    subjects = json.loads(row["unique_subjects"] or "[]")
    if subject and subject not in subjects:
        subjects.append(subject)

    total_points = int(row["total_points"] or 0) + int(points)
    quizzes_taken = int(row["quizzes_taken"] or 0) + int(bool(quiz))
    perfect_scores = int(row["perfect_scores"] or 0) + int(bool(perfect))
    level = total_points // 100 + 1

    conn.execute(
        """UPDATE user_stats SET total_points=?, level=?,
        quizzes_taken=?, perfect_scores=?, unique_subjects=?
        WHERE username=?""",
        (total_points, level, quizzes_taken, perfect_scores,
         json.dumps(subjects, ensure_ascii=False), u)
    )
    conn.commit()
    conn.close()
    check_badges(u)


def save_quiz_result(u, lid, title, subject, score, total):
    percent = round(100 * score / total, 2) if total else 0
    conn = get_db()
    conn.execute(
        """INSERT INTO quiz_history
        (username, lesson_id, lesson_title, subject, score, total, percent)
        VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (u, lid, title, subject, score, total, percent)
    )
    conn.commit()
    conn.close()
    points = score * 10
    update_user_stats(u, points, True, total > 0 and score == total, subject)
    return percent


def get_quiz_history(u, limit=20):
    conn = get_db()
    rows = conn.execute(
        """SELECT * FROM quiz_history WHERE username=?
        ORDER BY id DESC LIMIT ?""", (u, int(limit))
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_subject_stats(u):
    conn = get_db()
    rows = conn.execute(
        """SELECT subject, SUM(score) AS correct,
        SUM(total) AS total, COUNT(*) AS attempts,
        AVG(percent) AS avg_percent
        FROM quiz_history WHERE username=?
        GROUP BY subject ORDER BY attempts DESC""", (u,)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_weaknesses(u):
    stats = get_subject_stats(u)
    return sorted(
        [s for s in stats if s["avg_percent"] is not None],
        key=lambda s: s["avg_percent"]
    )


def get_leaderboard(period="all"):
    conn = get_db()
    where = ""
    if period == "week":
        where = "AND q.created_at >= datetime('now', '-7 days')"
    elif period == "month":
        where = "AND q.created_at >= datetime('now', '-30 days')"

    rows = conn.execute(f"""
        SELECT q.username, SUM(q.score * 10) AS points,
        COUNT(*) AS quizzes, AVG(q.percent) AS average
        FROM quiz_history q
        WHERE 1=1 {where}
        GROUP BY q.username
        ORDER BY points DESC, average DESC
        LIMIT 50
    """).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_weekly_report(u):
    conn = get_db()
    row = conn.execute(
        """SELECT COUNT(*) AS quizzes,
        COALESCE(SUM(score),0) AS correct,
        COALESCE(SUM(total),0) AS questions,
        COALESCE(AVG(percent),0) AS average
        FROM quiz_history
        WHERE username=? AND created_at >= datetime('now','-7 days')""",
        (u,)
    ).fetchone()
    conn.close()
    return dict(row)


def check_badges(u):
    s = get_user_stats(u)
    current = set(json.loads(s.get("badges") or "[]"))
    n = int(s.get("quizzes_taken") or 0)
    level = int(s.get("level") or 1)
    streak = int(s.get("streak") or 0)
    subjects = json.loads(s.get("unique_subjects") or "[]")

    conditions = {
        "first_quiz": n >= 1,
        "perfect": int(s.get("perfect_scores") or 0) >= 1,
        "5_quizzes": n >= 5,
        "10_quizzes": n >= 10,
        "25_quizzes": n >= 25,
        "50_quizzes": n >= 50,
        "100_quizzes": n >= 100,
        "level_5": level >= 5,
        "level_10": level >= 10,
        "level_20": level >= 20,
        "level_50": level >= 50,
        "all_subjects": len(subjects) >= len(SUBJECTS),
        "streak_7": streak >= 7,
        "streak_30": streak >= 30,
    }

    newly_earned = []
    for key, condition in conditions.items():
        if condition and key not in current:
            current.add(key)
            newly_earned.append(BADGES[key])
            add_notification(u, "إنجاز جديد", BADGES[key], "🏆")

    conn = get_db()
    conn.execute(
        "UPDATE user_stats SET badges=? WHERE username=?",
        (json.dumps(list(current)), u)
    )
    conn.commit()
    conn.close()
    return newly_earned


def get_user_badges(u):
    s = get_user_stats(u)
    return [BADGES[b] for b in json.loads(s.get("badges") or "[]")
            if b in BADGES]


def add_notification(u, title, msg, icon="🔔"):
    conn = get_db()
    conn.execute(
        """INSERT INTO notifications
        (username, title, message, icon) VALUES (?, ?, ?, ?)""",
        (u, title, msg, icon)
    )
    conn.commit()
    conn.close()


def get_notifications(u, unread=False):
    conn = get_db()
    if unread:
        rows = conn.execute(
            """SELECT * FROM notifications
            WHERE username=? AND is_read=0 ORDER BY id DESC""", (u,)
        ).fetchall()
    else:
        rows = conn.execute(
            """SELECT * FROM notifications
            WHERE username=? ORDER BY id DESC LIMIT 100""", (u,)
        ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def mark_notifications_read(u):
    conn = get_db()
    conn.execute(
        "UPDATE notifications SET is_read=1 WHERE username=?", (u,)
    )
    conn.commit()
    conn.close()


def check_daily_bonus(u):
    today = datetime.date.today().isoformat()
    conn = get_db()
    conn.execute(
        "INSERT OR IGNORE INTO user_stats (username) VALUES (?)", (u,)
    )
    row = conn.execute(
        "SELECT last_daily, streak FROM user_stats WHERE username=?", (u,)
    ).fetchone()

    if row["last_daily"] == today:
        conn.close()
        return 0

    yesterday = (
        datetime.date.today() - datetime.timedelta(days=1)
    ).isoformat()

    streak = int(row["streak"] or 0) + 1 if row["last_daily"] == yesterday else 1
    bonus = min(50, 10 + (streak - 1) * 5)

    conn.execute(
        """UPDATE user_stats SET last_daily=?, streak=?,
        total_points=total_points+?, level=(total_points+?) / 100 + 1
        WHERE username=?""",
        (today, streak, bonus, bonus, u)
    )
    conn.commit()
    conn.close()

    add_notification(u, "المكافأة اليومية", f"ربحت {bonus} نقطة!", "🎁")
    check_badges(u)
    return bonus


def get_rank(lvl):
    for minimum, title in RANKS:
        if lvl >= minimum:
            return title
    return "🌱 مبتدئ"


# =========================================================
# الدروس والأسئلة
# =========================================================

def load_lessons(subject=None, language=None, owner=None, search=""):
    conn = get_db()
    sql = "SELECT * FROM lessons WHERE 1=1"
    params = []

    if subject and subject != "الكل":
        sql += " AND subject=?"
        params.append(subject)
    if language and language != "الكل":
        sql += " AND language=?"
        params.append(language)
    if owner:
        sql += " AND owner=?"
        params.append(owner)
    if search.strip():
        sql += " AND (title LIKE ? OR content LIKE ?)"
        params.extend([f"%{search}%", f"%{search}%"])

    sql += " ORDER BY id DESC"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def add_lesson(subject, language, title, content, image_url="",
               pdf_url="", owner="soufianeDEV"):
    conn = get_db()
    cur = conn.execute(
        """INSERT INTO lessons
        (subject, language, title, content, image_url, pdf_url, owner)
        VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (subject, language, title, content, image_url, pdf_url, owner)
    )
    lid = cur.lastrowid
    conn.commit()
    conn.close()
    return lid


def delete_lesson(lid):
    conn = get_db()
    conn.execute("DELETE FROM questions WHERE lesson_id=?", (lid,))
    conn.execute("DELETE FROM lessons WHERE id=?", (lid,))
    conn.execute("DELETE FROM favorites WHERE lesson_id=?", (lid,))
    conn.execute("DELETE FROM reviews WHERE lesson_id=?", (lid,))
    conn.execute("DELETE FROM messages WHERE lesson_id=?", (lid,))
    conn.execute("DELETE FROM notes WHERE lesson_id=?", (lid,))
    conn.commit()
    conn.close()


def load_questions(lid):
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM questions WHERE lesson_id=? ORDER BY id", (lid,)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def add_question(lid, question, a, b, c, d, correct, explanation=""):
    if correct not in ["A", "B", "C", "D"]:
        raise ValueError("الإجابة الصحيحة يجب أن تكون A أو B أو C أو D.")
    conn = get_db()
    conn.execute(
        """INSERT INTO questions
        (lesson_id, question, option_a, option_b, option_c,
        option_d, correct_answer, explanation)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
        (lid, question, a, b, c, d, correct, explanation)
    )
    conn.commit()
    conn.close()


def delete_question(qid):
    conn = get_db()
    conn.execute("DELETE FROM questions WHERE id=?", (qid,))
    conn.commit()
    conn.close()


def upload_file(f, folder=UPLOAD_DIR):
    os.makedirs(folder, exist_ok=True)
    safe_name = re.sub(r"[^a-zA-Z0-9_.-]", "_", pathlib.Path(f.name).name)
    unique_name = f"{int(time.time())}_{random.randint(1000,9999)}_{safe_name}"
    path = os.path.join(folder, unique_name)
    with open(path, "wb") as output:
        output.write(f.getbuffer())
    return path


def extract_text_from_file(path):
    ext = pathlib.Path(path).suffix.lower()

    if ext == ".txt":
        for encoding in ("utf-8", "utf-8-sig", "cp1252"):
            try:
                return pathlib.Path(path).read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        return ""

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            return f"تعذر استخراج النص من PDF: {exc}"

    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(c.text for c in row.cells))
            return "\n".join(paragraphs)
        except Exception as exc:
            return f"تعذر استخراج النص من DOCX: {exc}"

    return ""


def render_pdf(path):
    if not path or not os.path.isfile(path):
        st.error("ملف PDF غير موجود.")
        return
    data = pathlib.Path(path).read_bytes()
    encoded = base64.b64encode(data).decode("utf-8")
    st.markdown(
        f'<iframe src="data:application/pdf;base64,{encoded}" '
        'width="100%" height="650" style="border:0"></iframe>',
        unsafe_allow_html=True
    )


# =========================================================
# توليد أسئلة من محتوى الدرس
# =========================================================

def clean_text_for_questions(content):
    content = re.sub(r"https?://\S+", "", content)
    content = re.sub(r"[ \t]+", " ", content)
    content = re.sub(r"\n+", "\n", content)
    return content.strip()


def generate_questions_from_content(content, count=10):
    """
    مولد اقتراحات محلي، لا يحتاج إلى API خارجي.
    يقترح أسئلة عن التعاريف والمعلومات الواردة في النص.
    يجب مراجعة الأسئلة والإجابات قبل اعتمادها.
    """
    content = clean_text_for_questions(content)
    if len(content) < 40:
        return []

    sentences = re.split(r"(?<=[.!؟?。])\s+|\n+", content)
    sentences = [
        re.sub(r"^[\-\*\d\.\s]+", "", s).strip()
        for s in sentences
    ]
    sentences = [
        s for s in sentences
        if 35 <= len(s) <= 350 and len(s.split()) >= 5
    ]

    # البحث عن تعريفات من نمط: X هو ... / X هي ...
    candidates = []
    for sentence in sentences:
        patterns = [
            r"^(.{2,70}?)\s+(?:هو|هي|يعني|تعني|يسمى|تسمى)\s+(.{10,200})$",
            r"^(.{2,70}?)\s+(?:est|sont|désigne|signifie|is|are|means)\s+(.{10,200})$",
        ]
        for pattern in patterns:
            match = re.match(pattern, sentence, flags=re.IGNORECASE)
            if match:
                term, definition = match.groups()
                candidates.append({
                    "question": f"ما المقصود بـ «{term.strip()}»؟",
                    "answer": sentence,
                    "source": sentence,
                    "term": term.strip(),
                    "definition": definition.strip()
                })
                break

    # الجمل التي تتضمن معلومة قابلة للاسترجاع.
    for sentence in sentences:
        if not any(c["source"] == sentence for c in candidates):
            candidates.append({
                "question": "أي عبارة تلخّص المعلومة التالية بشكل صحيح؟",
                "answer": sentence,
                "source": sentence,
                "term": "",
                "definition": ""
            })

    # إزالة التكرار
    unique = []
    seen = set()
    for candidate in candidates:
        key = candidate["answer"].casefold()
        if key not in seen:
            seen.add(key)
            unique.append(candidate)

    random.shuffle(unique)
    results = []

    for candidate in unique[:count]:
        answer = candidate["answer"]
        alternatives = [
            item["answer"] for item in unique
            if item["answer"] != answer
        ]
        random.shuffle(alternatives)
        options = [answer] + alternatives[:3]

        # إذا لم يتوفر إلا جواب واحد، نضيف مشتتات صريحة.
        while len(options) < 4:
            options.append("لا توجد معلومة تؤكد هذه العبارة في النص.")

        random.shuffle(options)
        correct_letter = ["A", "B", "C", "D"][options.index(answer)]

        # تنبيه: هذه الأسئلة تقيس مطابقة النص، وليست فهماً دلالياً كاملاً.
        results.append({
            "question": candidate["question"],
            "options": options,
            "correct": correct_letter,
            "explanation": "الجواب مستخرج من محتوى الدرس: " + answer
        })

    return results


def save_generated_questions(lid, generated):
    for q in generated:
        add_question(
            lid,
            q["question"],
            q["options"][0],
            q["options"][1],
            q["options"][2],
            q["options"][3],
            q["correct"],
            q["explanation"]
        )


# =========================================================
# المفضلة والمراجعات والنقاش والملاحظات
# =========================================================

def toggle_favorite(u, lid):
    conn = get_db()
    row = conn.execute(
        "SELECT id FROM favorites WHERE username=? AND lesson_id=?",
        (u, lid)
    ).fetchone()
    if row:
        conn.execute("DELETE FROM favorites WHERE id=?", (row["id"],))
    else:
        conn.execute(
            "INSERT OR IGNORE INTO favorites (username, lesson_id) VALUES (?, ?)",
            (u, lid)
        )
    conn.commit()
    conn.close()


def is_favorite(u, lid):
    conn = get_db()
    row = conn.execute(
        "SELECT 1 FROM favorites WHERE username=? AND lesson_id=?",
        (u, lid)
    ).fetchone()
    conn.close()
    return bool(row)


def get_favorites(u):
    conn = get_db()
    rows = conn.execute(
        """SELECT l.* FROM lessons l
        JOIN favorites f ON l.id=f.lesson_id
        WHERE f.username=? ORDER BY f.id DESC""", (u,)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def add_review(lid, u, rating, comment):
    conn = get_db()
    conn.execute(
        """INSERT INTO reviews (lesson_id, username, rating, comment)
        VALUES (?, ?, ?, ?)
        ON CONFLICT(lesson_id, username) DO UPDATE SET
        rating=excluded.rating, comment=excluded.comment""",
        (lid, u, int(rating), comment)
    )
    conn.commit()
    conn.close()


def get_reviews(lid):
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM reviews WHERE lesson_id=? ORDER BY rowid DESC", (lid,)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_avg_rating(lid):
    conn = get_db()
    row = conn.execute(
        "SELECT AVG(rating) AS rating FROM reviews WHERE lesson_id=?", (lid,)
    ).fetchone()
    conn.close()
    return round(row["rating"] or 0, 1)


def add_message(lid, u, message):
    conn = get_db()
    conn.execute(
        "INSERT INTO messages (lesson_id, username, message) VALUES (?, ?, ?)",
        (lid, u, message)
    )
    conn.commit()
    conn.close()


def get_messages(lid):
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM messages WHERE lesson_id=? ORDER BY id DESC LIMIT 100",
        (lid,)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def add_note(u, lid, note):
    conn = get_db()
    conn.execute(
        "INSERT INTO notes (username, lesson_id, note) VALUES (?, ?, ?)",
        (u, lid, note)
    )
    conn.commit()
    conn.close()


def get_notes(u, lid):
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM notes WHERE username=? AND lesson_id=? ORDER BY id DESC",
        (u, lid)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def delete_note(nid):
    conn = get_db()
    conn.execute(
        "DELETE FROM notes WHERE id=? AND username=?",
        (nid, st.session_state.get("username", ""))
    )
    conn.commit()
    conn.close()


# =========================================================
# بطاقات الحفظ
# =========================================================

def add_flashcard(u, subject, front, back):
    conn = get_db()
    conn.execute(
        "INSERT INTO flashcards (username, subject, front, back) VALUES (?, ?, ?, ?)",
        (u, subject, front, back)
    )
    conn.commit()
    conn.close()


def get_flashcards(u, s=None):
    conn = get_db()
    if s and s != "الكل":
        rows = conn.execute(
            "SELECT * FROM flashcards WHERE username=? AND subject=? ORDER BY id DESC",
            (u, s)
        ).fetchall()
    else:
        rows = conn.execute(
            "SELECT * FROM flashcards WHERE username=? ORDER BY id DESC", (u,)
        ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def delete_flashcard(fid):
    conn = get_db()
    conn.execute(
        "DELETE FROM flashcards WHERE id=? AND username=?",
        (fid, st.session_state.username)
    )
    conn.commit()
    conn.close()


def toggle_flashcard_known(fid):
    conn = get_db()
    conn.execute(
        """UPDATE flashcards SET known=1-known
        WHERE id=? AND username=?""",
        (fid, st.session_state.username)
    )
    conn.commit()
    conn.close()


# =========================================================
# خطة الدراسة والأصدقاء
# =========================================================

def add_friend(u, fu):
    if u == fu:
        return False
    conn = get_db()
    exists = conn.execute(
        "SELECT 1 FROM users WHERE username=?", (fu,)
    ).fetchone()
    if not exists:
        conn.close()
        return False
    conn.execute(
        """INSERT OR IGNORE INTO friends (username, friend_username, status)
        VALUES (?, ?, 'accepted')""", (u, fu)
    )
    conn.commit()
    conn.close()
    return True


def get_friends(u):
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM friends WHERE username=? ORDER BY id DESC", (u,)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def remove_friend(u, fu):
    conn = get_db()
    conn.execute(
        "DELETE FROM friends WHERE username=? AND friend_username=?", (u, fu)
    )
    conn.commit()
    conn.close()


def add_study_plan(u, subject, priority, target_date):
    conn = get_db()
    conn.execute(
        """INSERT INTO study_plan
        (username, subject, priority, target_date)
        VALUES (?, ?, ?, ?)""",
        (u, subject, priority, str(target_date))
    )
    conn.commit()
    conn.close()


def get_study_plan(u):
    conn = get_db()
    rows = conn.execute(
        """SELECT * FROM study_plan WHERE username=?
        ORDER BY completed, target_date""", (u,)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def toggle_study_plan(pid):
    conn = get_db()
    conn.execute(
        """UPDATE study_plan SET completed=1-completed
        WHERE id=? AND username=?""",
        (pid, st.session_state.username)
    )
    conn.commit()
    conn.close()


def auto_generate_plan(u):
    existing = get_study_plan(u)
    if existing:
        return
    today = datetime.date.today()
    for i, subject in enumerate(SUBJECTS.keys()):
        add_study_plan(
            u, subject, "متوسطة",
            today + datetime.timedelta(days=i)
        )


# =========================================================
# الثيم والتنسيق
# =========================================================

def apply_theme():
    theme_name = st.session_state.get("theme", "🌙 Moonlight")
    theme = THEMES.get(theme_name, THEMES["🌙 Moonlight"])

    st.markdown(f"""
    <style>
    :root {{
        --bg: {theme['bg']};
        --card: {theme['card']};
        --text: {theme['text']};
        --accent: {theme['accent']};
        --secondary: {theme['secondary']};
        --border: {theme['border']};
        --highlight: {theme['highlight']};
    }}
    html, body, [data-testid="stAppViewContainer"] {{
        direction: rtl;
        text-align: right;
        background-color: var(--bg);
        color: var(--text);
    }}
    [data-testid="stSidebar"] {{
        direction: rtl;
        text-align: right;
        background-color: var(--secondary);
    }}
    [data-testid="stHeader"] {{
        background-color: var(--bg);
    }}
    .stMarkdown, label, p, h1, h2, h3, h4, li {{
        text-align: right;
        color: var(--text);
    }}
    [data-testid="stVerticalBlockBorderWrapper"] {{
        border-color: var(--border);
    }}
    .hero {{
        background: linear-gradient(120deg, var(--card), var(--secondary));
        padding: 28px;
        border-radius: 20px;
        border: 1px solid var(--border);
        margin-bottom: 20px;
    }}
    .hero h1, .hero p {{
        color: var(--text);
    }}
    .lesson-card {{
        background: var(--card);
        border: 1px solid var(--border);
        border-radius: 16px;
        padding: 18px;
        margin: 10px 0;
    }}
    .metric-card {{
        background: var(--card);
        border: 1px solid var(--border);
        border-radius: 14px;
        padding: 16px;
        text-align: center;
    }}
    .footer {{
        text-align: center;
        padding: 22px 5px;
        margin-top: 30px;
        border-top: 1px solid var(--border);
        opacity: .85;
        direction: ltr;
    }}
    div.stButton > button {{
        border-radius: 10px;
        min-height: 42px;
    }}
    input, textarea {{
        text-align: right !important;
    }}
    [data-testid="stMetricValue"] {{
        direction: ltr;
    }}
    </style>
    """, unsafe_allow_html=True)


def render_footer():
    st.markdown(
        '<div class="footer">© 2026 Soufiane Ouhazza — 3AC RevisioMaroc</div>',
        unsafe_allow_html=True
    )


def subject_label(key):
    return SUBJECTS.get(key, {}).get("name", key)


def get_lesson_by_id(lid):
    conn = get_db()
    row = conn.execute("SELECT * FROM lessons WHERE id=?", (lid,)).fetchone()
    conn.close()
    return dict(row) if row else None


# =========================================================
# المصادقة
# =========================================================

def login_user(user):
    st.session_state.authenticated = True
    st.session_state.username = user["username"]
    st.session_state.role = user["role"]
    st.session_state.full_name = user["full_name"] or user["username"]
    st.session_state.page = "dashboard"
    st.session_state.pop("quiz_run", None)

    bonus = check_daily_bonus(user["username"])
    if bonus:
        st.session_state.login_bonus = bonus
    else:
        st.session_state.login_bonus = 0


def render_auth_home():
    st.markdown("""
    <div class="hero">
        <h1>📚 3AC RevisioMaroc</h1>
        <p>منصتك لمراجعة دروس الثالثة إعدادي، إنشاء الاختبارات وتتبع تقدمك.</p>
    </div>
    """, unsafe_allow_html=True)

    left, center, right = st.columns(3)
    with left:
        st.markdown("### 🎓")
        st.subheader("فضاء التلميذ")
        st.write("راجع الدروس، حل الاختبارات واجمع النقاط.")
        if st.button("الدخول كتلميذ", use_container_width=True):
            st.session_state.auth_page = "student"
            st.rerun()
    with center:
        st.markdown("### 🛠️")
        st.subheader("فضاء المطور")
        st.write("أضف الدروس والأسئلة وأدر المحتوى.")
        if st.button("الدخول كمطور", use_container_width=True):
            st.session_state.auth_page = "developer"
            st.rerun()
    with right:
        st.markdown("### ✍️")
        st.subheader("حساب جديد")
        st.write("أنشئ حساباً وابدأ رحلة المراجعة.")
        if st.button("إنشاء حساب", use_container_width=True):
            st.session_state.auth_page = "register"
            st.rerun()


def render_student_login():
    st.subheader("🎓 تسجيل دخول التلميذ")
    with st.form("student_login"):
        username = st.text_input("اسم المستخدم")
        password = st.text_input("كلمة المرور", type="password")
        submitted = st.form_submit_button("دخول", use_container_width=True)
    if submitted:
        user = authenticate(username, password)
        if user and user["role"] == "student":
            login_user(user)
            st.rerun()
        st.error("معلومات الدخول غير صحيحة أو الحساب ليس حساب تلميذ.")
    if st.button("رجوع"):
        st.session_state.auth_page = "home"
        st.rerun()


def render_developer_login():
    st.subheader("🛠️ تسجيل دخول المطور")
    st.info("بيانات المطور الافتراضية موضحة في إعدادات التطبيق. غيّرها قبل نشر التطبيق.")
    with st.form("developer_login"):
        username = st.text_input("اسم المستخدم")
        password = st.text_input("كلمة المرور", type="password")
        submitted = st.form_submit_button("دخول المطور", use_container_width=True)
    if submitted:
        user = authenticate(username, password)
        if user and user["role"] == "developer":
            login_user(user)
            st.rerun()
        st.error("بيانات المطور غير صحيحة.")
    if st.button("رجوع"):
        st.session_state.auth_page = "home"
        st.rerun()


def render_register():
    st.subheader("✍️ إنشاء حساب")
    with st.form("register_form"):
        name = st.text_input("الاسم الكامل")
        username = st.text_input("اسم المستخدم")
        password = st.text_input("كلمة المرور", type="password")
        password2 = st.text_input("تأكيد كلمة المرور", type="password")
        submitted = st.form_submit_button("إنشاء الحساب", use_container_width=True)
    if submitted:
        if password != password2:
            st.error("كلمتا المرور غير متطابقتين.")
        else:
            ok, message = register_user(username, password, name)
            if ok:
                st.success(message)
                st.session_state.auth_page = "student"
                st.rerun()
            else:
                st.error(message)
    if st.button("رجوع"):
        st.session_state.auth_page = "home"
        st.rerun()


def render_auth_page():
    page = st.session_state.get("auth_page", "home")
    if page == "home":
        render_auth_home()
    elif page == "student":
        render_student_login()
    elif page == "developer":
        render_developer_login()
    elif page == "register":
        render_register()
    render_footer()


# =========================================================
# الصفحة الرئيسية
# =========================================================

def render_dashboard():
    u = st.session_state.username
    s = get_user_stats(u)

    st.markdown(f"""
    <div class="hero">
        <h1>مرحبا {html.escape(st.session_state.full_name)} 👋</h1>
        <p>أهلاً بك في 3AC RevisioMaroc. كل يوم فرصة جديدة للتفوق.</p>
    </div>
    """, unsafe_allow_html=True)

    bonus = st.session_state.pop("login_bonus", 0)
    if bonus:
        st.success(f"🎁 المكافأة اليومية: +{bonus} نقطة")

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("⭐ مجموع النقاط", s["total_points"])
    c2.metric("🏅 المستوى", s["level"])
    c3.metric("📝 الاختبارات", s["quizzes_taken"])
    c4.metric("💯 علامات كاملة", s["perfect_scores"])

    st.info(f"لقبك الحالي: {get_rank(s['level'])}")

    st.subheader("📚 اختر المادة")
    cols = st.columns(4)
    for i, (key, subject) in enumerate(SUBJECTS.items()):
        with cols[i % 4]:
            st.markdown(
                f"""<div class="lesson-card">
                <h2>{subject['icon']}</h2>
                <h4>{subject['name']}</h4></div>""",
                unsafe_allow_html=True
            )
            if st.button("عرض الدروس", key=f"subject_{key}",
                         use_container_width=True):
                st.session_state.selected_subject = key
                st.session_state.page = "lessons"
                st.rerun()

    st.subheader("📈 آخر النتائج")
    history = get_quiz_history(u, 5)
    if history:
        for item in history:
            st.write(
                f"• {item['lesson_title']} — "
                f"{item['score']}/{item['total']} "
                f"({item['percent']}%)"
            )
    else:
        st.write("لم تجتز أي اختبار بعد. ابدأ باختيار درس.")

    st.subheader("🎯 إنجازاتك")
    badges = get_user_badges(u)
    st.write(" • ".join(badges) if badges else "واصل التقدم لتحصل على أول شارة!")


# =========================================================
# الدروس
# =========================================================

def _render_reviews(lid):
    st.subheader("⭐ تقييم الدرس")
    reviews = get_reviews(lid)
    st.write(f"متوسط التقييم: {get_avg_rating(lid)} / 5")
    for review in reviews:
        st.write(
            f"⭐ {review['rating']}/5 — "
            f"**{review['username']}**: {review['comment']}"
        )

    with st.form(f"review_{lid}"):
        rating = st.slider("تقييمك", 1, 5, 5)
        comment = st.text_area("تعليقك")
        submit = st.form_submit_button("حفظ التقييم")
    if submit:
        add_review(lid, st.session_state.username, rating, comment)
        st.success("تم حفظ التقييم.")
        st.rerun()


def _render_discussion(lid):
    st.subheader("💬 المناقشة")
    for message in get_messages(lid):
        st.markdown(
            f"**{html.escape(message['username'])}:** "
            f"{html.escape(message['message'])}"
        )
    with st.form(f"message_{lid}"):
        message = st.text_area("اكتب تعليقاً أو سؤالاً")
        submit = st.form_submit_button("إرسال")
    if submit:
        if message.strip():
            add_message(lid, st.session_state.username, message.strip())
            st.rerun()
        else:
            st.warning("اكتب رسالة أولاً.")


def _render_notes(lid):
    st.subheader("🗒️ ملاحظاتي")
    for note in get_notes(st.session_state.username, lid):
        col1, col2 = st.columns([5, 1])
        with col1:
            st.write(note["note"])
        with col2:
            if st.button("حذف", key=f"delnote_{note['id']}"):
                delete_note(note["id"])
                st.rerun()

    with st.form(f"note_{lid}"):
        note_text = st.text_area("أضف ملاحظة شخصية")
        submit = st.form_submit_button("حفظ الملاحظة")
    if submit:
        if note_text.strip():
            add_note(st.session_state.username, lid, note_text.strip())
            st.success("تم حفظ الملاحظة.")
            st.rerun()


def _render_lesson_card(lesson):
    lid = lesson["id"]
    st.markdown(
        f"""<div class="lesson-card">
        <h3>{SUBJECTS.get(lesson['subject'], {}).get('icon', '📚')}
        {html.escape(lesson['title'])}</h3>
        <p>{subject_label(lesson['subject'])}</p>
        </div>""",
        unsafe_allow_html=True
    )

    b1, b2, b3 = st.columns(3)
    with b1:
        if st.button("⭐ المفضلة", key=f"fav_{lid}"):
            toggle_favorite(st.session_state.username, lid)
            st.rerun()
    with b2:
        if st.button("📖 فتح الدرس", key=f"open_{lid}"):
            st.session_state.current_lesson_id = lid
            st.session_state.page = "lesson_detail"
            st.rerun()
    with b3:
        if st.button("📝 اختبار", key=f"quizlesson_{lid}"):
            st.session_state.quiz_lesson_ids = [lid]
            st.session_state.page = "quiz"
            st.rerun()


def render_lessons():
    st.title("📚 مكتبة الدروس")

    c1, c2, c3 = st.columns([2, 1, 1])
    with c1:
        search = st.text_input("🔎 ابحث عن درس أو كلمة")
    with c2:
        subject_options = ["الكل"] + list(SUBJECTS.keys())
        subject = st.selectbox(
            "المادة",
            subject_options,
            format_func=lambda x: "جميع المواد" if x == "الكل" else subject_label(x),
            index=(subject_options.index(st.session_state.get("selected_subject"))
                   if st.session_state.get("selected_subject") in subject_options else 0)
        )
    with c3:
        language = st.selectbox(
            "لغة الدرس", ["الكل", "ar", "fr", "en"],
            format_func=lambda x: {
                "الكل": "جميع اللغات", "ar": "العربية",
                "fr": "الفرنسية", "en": "الإنجليزية"
            }[x]
        )

    lessons = load_lessons(subject, language, search=search)
    st.caption(f"عدد الدروس: {len(lessons)}")

    if not lessons:
        st.info("لا توجد دروس مطابقة. يمكن للمطور إضافة دروس جديدة.")
    for lesson in lessons:
        _render_lesson_card(lesson)


def render_lesson_detail():
    lid = st.session_state.get("current_lesson_id")
    lesson = get_lesson_by_id(lid) if lid else None
    if not lesson:
        st.error("الدرس غير موجود.")
        return

    st.title(lesson["title"])
    st.caption(subject_label(lesson["subject"]))

    tab1, tab2, tab3, tab4 = st.tabs([
        "📖 محتوى الدرس", "⭐ التقييم", "💬 النقاش", "🗒️ ملاحظاتي"
    ])

    with tab1:
        if lesson.get("image_url"):
            st.image(lesson["image_url"], use_container_width=True)
        st.markdown(lesson.get("content") or "لا يوجد نص للدرس.")
        path = lesson.get("pdf_url", "")
        if path and os.path.isfile(path):
            st.subheader("📄 ملف الدرس")
            render_pdf(path)
            with open(path, "rb") as f:
                st.download_button(
                    "⬇️ تحميل ملف الدرس",
                    data=f.read(),
                    file_name=os.path.basename(path),
                    mime="application/pdf" if path.lower().endswith(".pdf")
                    else "application/octet-stream"
                )

        if st.button("📝 ابدأ اختبار هذا الدرس"):
            st.session_state.quiz_lesson_ids = [lid]
            st.session_state.page = "quiz"
            st.rerun()

    with tab2:
        _render_reviews(lid)
    with tab3:
        _render_discussion(lid)
    with tab4:
        _render_notes(lid)


# =========================================================
# الاختبار متعدد المواد مع اختيار الأسئلة والمؤقت
# =========================================================

def collect_quiz_questions(lesson_ids, selected_subjects):
    all_questions = []
    conn = get_db()

    for lid in lesson_ids:
        lesson = conn.execute(
            "SELECT * FROM lessons WHERE id=?", (lid,)
        ).fetchone()
        if not lesson:
            continue
        if selected_subjects and lesson["subject"] not in selected_subjects:
            continue

        rows = conn.execute(
            "SELECT * FROM questions WHERE lesson_id=? ORDER BY id", (lid,)
        ).fetchall()
        for row in rows:
            q = dict(row)
            q["_lesson_id"] = lid
            q["_lesson_title"] = lesson["title"]
            q["_subject"] = lesson["subject"]
            all_questions.append(q)

    conn.close()
    return all_questions


def start_quiz_session(questions, duration, shuffle_questions=True):
    if shuffle_questions:
        random.shuffle(questions)
    now = time.time()
    st.session_state.quiz_run = {
        "questions": questions,
        "answers": {},
        "started_at": now,
        "deadline": now + duration,
        "duration": duration,
        "submitted": False,
        "result": None,
    }
    st.session_state.quiz_widget_seed = random.randint(10000, 99999999)


def grade_quiz():
    run = st.session_state.quiz_run
    score = 0
    for index, q in enumerate(run["questions"]):
        selected = run["answers"].get(str(index))
        if selected == q["correct_answer"]:
            score += 1

    total = len(run["questions"])
    results = {
        "score": score,
        "total": total,
        "percent": round(score * 100 / total, 2) if total else 0,
    }
    run["submitted"] = True
    run["result"] = results

    # حفظ النتيجة كاختبار واحد، مع اعتماد المادة إذا كانت موحدة.
    subjects = {q["_subject"] for q in run["questions"]}
    subject = list(subjects)[0] if len(subjects) == 1 else "mixed"
    lesson_ids = list({q["_lesson_id"] for q in run["questions"]})
    lesson_title = (
        get_lesson_by_id(lesson_ids[0])["title"]
        if len(lesson_ids) == 1 and get_lesson_by_id(lesson_ids[0])
        else "اختبار متعدد الدروس"
    )
    save_quiz_result(
        st.session_state.username,
        lesson_ids[0] if len(lesson_ids) == 1 else None,
        lesson_title, subject, score, total
    )
    return results


def _render_quiz_ui(questions, lesson_id=None):
    run = st.session_state.get("quiz_run")
    if not run:
        st.info("اختر الدروس ثم اضغط على بدء الاختبار.")
        return

    if run["submitted"]:
        result = run["result"]
        st.success(
            f"النتيجة: {result['score']} / {result['total']} "
            f"— {result['percent']}%"
        )
        st.progress(result["percent"] / 100)

        for i, q in enumerate(run["questions"]):
            selected = run["answers"].get(str(i), "لم تجب")
            is_correct = selected == q["correct_answer"]
            st.markdown(f"### السؤال {i+1}: {q['question']}")
            st.write(f"إجابتك: {selected}")
            st.write(f"الإجابة الصحيحة: {q['correct_answer']}")
            if is_correct:
                st.success("إجابة صحيحة")
            else:
                st.error("إجابة غير صحيحة")
            if q.get("explanation"):
                st.info(q["explanation"])

        if st.button("🔄 اختبار جديد"):
            st.session_state.pop("quiz_run", None)
            st.rerun()
        return

    remaining = max(0, int(run["deadline"] - time.time()))
    minutes, seconds = divmod(remaining, 60)
    st.subheader(f"⏱️ الوقت المتبقي: {minutes:02d}:{seconds:02d}")
    st.progress(min(1.0, remaining / max(1, run["duration"])))

    if remaining <= 0:
        grade_quiz()
        st.warning("انتهى الوقت. تم تصحيح الاختبار.")
        st.rerun()

    st.caption(
        f"عدد الأسئلة: {len(run['questions'])} — "
        f"النقاط: 10 لكل إجابة صحيحة"
    )

    # منع الاعتماد على اختيار عشوائي من المستخدم من دون حفظ الإجابة.
    with st.form(f"quiz_form_{st.session_state.quiz_widget_seed}"):
        for i, q in enumerate(run["questions"]):
            st.markdown(f"### {i+1}. {q['question']}")
            options = {
                "A": q["option_a"],
                "B": q["option_b"],
                "C": q["option_c"],
                "D": q["option_d"],
            }
            current = run["answers"].get(str(i), "")
            choice = st.radio(
                "اختر جواباً:",
                ["A", "B", "C", "D"],
                index=["A", "B", "C", "D"].index(current) if current in options else None,
                format_func=lambda x, opts=options: f"{x}. {opts[x]}",
                key=f"answer_{st.session_state.quiz_widget_seed}_{i}"
            )
            # يُجمع الاختيار عند إرسال النموذج.
            st.session_state[f"_pending_answer_{i}"] = choice

        submitted = st.form_submit_button("✅ إنهاء وتصحيح الاختبار")

    if submitted:
        for i in range(len(run["questions"])):
            choice = st.session_state.get(f"_pending_answer_{i}")
            if choice in ("A", "B", "C", "D"):
                run["answers"][str(i)] = choice
        grade_quiz()
        st.rerun()

    st.caption("يُحتسب الوقت عند فتح صفحة الاختبار. اضغط على إنهاء لتصحيح إجاباتك.")


def render_quiz():
    st.title("📝 الاختبارات")

    lessons = load_lessons()
    if not lessons:
        st.warning("أضف دروساً وأسئلة أولاً.")
        return

    if "quiz_run" in st.session_state and st.session_state.quiz_run:
        _render_quiz_ui(
            st.session_state.quiz_run["questions"],
            None
        )
        return

    subject_keys = list(SUBJECTS.keys())
    selected_subjects = st.multiselect(
        "🎯 اختر مادة واحدة أو أكثر",
        subject_keys,
        default=subject_keys,
        format_func=subject_label
    )

    filtered = [
        l for l in lessons
        if not selected_subjects or l["subject"] in selected_subjects
    ]

    lesson_ids = st.multiselect(
        "📚 اختر الدروس التي تريد الاختبار فيها",
        [l["id"] for l in filtered],
        default=[
            lid for lid in st.session_state.get("quiz_lesson_ids", [])
            if any(l["id"] == lid for l in filtered)
        ],
        format_func=lambda lid: next(
            (l["title"] for l in filtered if l["id"] == lid), str(lid)
        )
    )

    all_questions = collect_quiz_questions(lesson_ids, selected_subjects)
    st.info(f"الأسئلة المتوفرة في اختياراتك: {len(all_questions)}")

    count = st.number_input(
        "🔢 عدد الأسئلة",
        min_value=1,
        max_value=max(1, min(100, len(all_questions))),
        value=max(1, min(10, len(all_questions))),
        step=1
    )
    duration_minutes = st.selectbox(
        "⏱️ مدة الاختبار",
        [5, 10, 15, 20, 30, 45, 60],
        index=1,
        format_func=lambda x: f"{x} دقائق"
    )
    shuffle_questions = st.checkbox("ترتيب الأسئلة عشوائياً", value=True)

    if st.button("🚀 بدء الاختبار", use_container_width=True):
        if not lesson_ids:
            st.warning("اختر درساً واحداً على الأقل.")
        elif not all_questions:
            st.warning("الدروس المحددة لا تحتوي على أسئلة.")
        else:
            chosen = all_questions[:]
            random.shuffle(chosen)
            chosen = chosen[:int(count)]
            start_quiz_session(
                chosen, int(duration_minutes) * 60, shuffle_questions
            )
            st.rerun()


# =========================================================
# المراجعة السريعة
# =========================================================

def render_quick_review():
    st.title("⚡ المراجعة السريعة")
    subject = st.selectbox(
        "المادة", ["الكل"] + list(SUBJECTS.keys()),
        format_func=lambda x: "كل المواد" if x == "الكل" else subject_label(x)
    )
    lessons = load_lessons(subject=subject)
    if not lessons:
        st.info("لا توجد دروس للمراجعة.")
        return

    for lesson in lessons:
        with st.expander(lesson["title"]):
            st.write(lesson["content"] or "لا يوجد محتوى.")
            questions = load_questions(lesson["id"])
            st.write(f"عدد الأسئلة: {len(questions)}")
            if st.button("اختبر نفسي", key=f"review_{lesson['id']}"):
                st.session_state.quiz_lesson_ids = [lesson["id"]]
                st.session_state.page = "quiz"
                st.rerun()


# =========================================================
# بطاقات الحفظ
# =========================================================

def render_flashcards():
    st.title("🗂️ بطاقات الحفظ")
    u = st.session_state.username

    with st.expander("➕ إنشاء بطاقة جديدة", expanded=False):
        with st.form("new_flashcard"):
            subject = st.selectbox(
                "المادة", list(SUBJECTS.keys()),
                format_func=subject_label
            )
            front = st.text_area("الوجه الأمامي: السؤال أو المصطلح")
            back = st.text_area("الوجه الخلفي: الجواب أو التعريف")
            submit = st.form_submit_button("حفظ البطاقة")
        if submit:
            if front.strip() and back.strip():
                add_flashcard(u, subject, front.strip(), back.strip())
                st.success("تمت إضافة البطاقة.")
                st.rerun()
            else:
                st.warning("أكمل السؤال والجواب.")

    subject_filter = st.selectbox(
        "تصفية حسب المادة", ["الكل"] + list(SUBJECTS.keys()),
        format_func=lambda x: "كل المواد" if x == "الكل" else subject_label(x)
    )
    cards = get_flashcards(u, subject_filter)
    if not cards:
        st.info("لا توجد بطاقات بعد.")
        return

    for card in cards:
        with st.container(border=True):
            st.caption(subject_label(card["subject"]))
            st.markdown(f"**السؤال:** {card['front']}")
            if st.toggle("إظهار الجواب", key=f"showcard_{card['id']}"):
                st.success(card["back"])
            c1, c2 = st.columns(2)
            with c1:
                label = "✅ أعرفها" if not card["known"] else "↩️ أراجعها"
                if st.button(label, key=f"known_{card['id']}"):
                    toggle_flashcard_known(card["id"])
                    st.rerun()
            with c2:
                if st.button("🗑️ حذف", key=f"deletecard_{card['id']}"):
                    delete_flashcard(card["id"])
                    st.rerun()


# =========================================================
# خطة الدراسة
# =========================================================

def render_study_plan():
    st.title("🗓️ خطة الدراسة")
    u = st.session_state.username

    with st.expander("➕ إضافة هدف دراسي"):
        with st.form("study_plan_form"):
            subject = st.selectbox(
                "المادة", list(SUBJECTS.keys()),
                format_func=subject_label
            )
            priority = st.selectbox("الأولوية", ["عالية", "متوسطة", "منخفضة"])
            target_date = st.date_input("تاريخ الإنجاز", value=datetime.date.today())
            submit = st.form_submit_button("إضافة الهدف")
        if submit:
            add_study_plan(u, subject, priority, target_date)
            st.rerun()

    if st.button("✨ إنشاء خطة تلقائية للمواد"):
        auto_generate_plan(u)
        st.success("تم إنشاء الخطة إذا لم تكن لديك أهداف من قبل.")
        st.rerun()

    plans = get_study_plan(u)
    if not plans:
        st.info("أضف أهدافك الدراسية لتنظيم وقتك.")
        return

    for plan in plans:
        cols = st.columns([1, 4, 2, 1])
        with cols[0]:
            checked = st.checkbox(
                "منجز",
                value=bool(plan["completed"]),
                key=f"plan_{plan['id']}"
            )
            # تُحفظ حالة التغيير دون قلبها عند كل إعادة عرض.
            if checked != bool(plan["completed"]):
                toggle_study_plan(plan["id"])
                st.rerun()
        with cols[1]:
            st.write(subject_label(plan["subject"]))
        with cols[2]:
            st.caption(f"{plan['priority']} — {plan['target_date']}")
        with cols[3]:
            st.write("✅" if plan["completed"] else "⏳")


# =========================================================
# الأصدقاء ولوحة المتصدرين والتقارير
# =========================================================

def render_friends():
    st.title("👥 زملائي")
    u = st.session_state.username

    with st.form("add_friend_form"):
        friend = st.text_input("اسم المستخدم للزميل")
        submit = st.form_submit_button("إضافة")
    if submit:
        if add_friend(u, friend.strip()):
            st.success("تمت إضافة الزميل.")
            st.rerun()
        else:
            st.error("تعذر إضافة الزميل. تأكد من اسم المستخدم.")

    for item in get_friends(u):
        col1, col2 = st.columns([4, 1])
        with col1:
            st.write(item["friend_username"])
        with col2:
            if st.button("حذف", key=f"remove_friend_{item['id']}"):
                remove_friend(u, item["friend_username"])
                st.rerun()


def render_leaderboard():
    st.title("🏆 لوحة المتصدرين")
    period = st.selectbox(
        "الفترة", ["all", "week", "month"],
        format_func=lambda x: {
            "all": "كل الوقت", "week": "آخر 7 أيام",
            "month": "آخر 30 يوماً"
        }[x]
    )
    rows = get_leaderboard(period)
    if not rows:
        st.info("لا توجد نتائج كافية حتى الآن.")
        return
    for i, row in enumerate(rows, 1):
        st.markdown(
            f"**{i}. {row['username']}** — "
            f"⭐ {row['points'] or 0} نقطة — "
            f"📝 {row['quizzes']} اختبار"
        )


def render_weekly_report():
    st.title("📊 التقرير الأسبوعي")
    report = get_weekly_report(st.session_state.username)
    a, b, c = st.columns(3)
    a.metric("الاختبارات", report["quizzes"])
    b.metric("الإجابات الصحيحة", report["correct"])
    c.metric("متوسط النتائج", f"{report['average']:.1f}%")
    st.write(
        f"عدد الأسئلة التي أجبت عنها: {report['questions']}"
    )

    weaknesses = get_weaknesses(st.session_state.username)
    st.subheader("🎯 المواد التي تحتاج إلى مراجعة")
    if weaknesses:
        for item in weaknesses:
            st.write(
                f"• {subject_label(item['subject'])}: "
                f"{item['avg_percent']:.1f}%"
            )
    else:
        st.info("أكمل بعض الاختبارات لعرض نقاط القوة والضعف.")


def render_notifications():
    st.title("🔔 الإشعارات")
    u = st.session_state.username
    notifications = get_notifications(u)

    if st.button("تعليم جميع الإشعارات كمقروءة"):
        mark_notifications_read(u)
        st.rerun()

    if not notifications:
        st.info("لا توجد إشعارات.")
        return

    for notification in notifications:
        with st.container(border=True):
            st.write(
                f"{notification['icon']} **{notification['title']}**"
            )
            st.write(notification["message"])
            if not notification["is_read"]:
                st.caption("جديد")


def render_my_stats():
    st.title("📈 إحصائياتي")
    u = st.session_state.username
    s = get_user_stats(u)
    history = get_quiz_history(u, 1000)
    subjects = get_subject_stats(u)

    st.metric("النقاط", s["total_points"])
    st.metric("المستوى", s["level"])
    st.metric("اللقب", get_rank(s["level"]))
    st.metric("سلسلة الأيام", s["streak"])

    st.subheader("🏅 الشارات")
    badges = get_user_badges(u)
    if badges:
        st.write("　".join(badges))
    else:
        st.info("لا توجد شارات بعد. ابدأ الاختبارات للحصول عليها.")

    st.subheader("📚 الأداء حسب المادة")
    if subjects:
        for item in subjects:
            st.write(
                f"**{subject_label(item['subject'])}** — "
                f"{item['avg_percent']:.1f}% — "
                f"{item['attempts']} اختبار"
            )

    st.subheader("📝 سجل الاختبارات")
    for item in history[:30]:
        st.write(
            f"{item['created_at']} | {item['lesson_title']} | "
            f"{item['score']}/{item['total']} | {item['percent']}%"
        )


def render_favorites():
    st.title("⭐ دروسي المفضلة")
    lessons = get_favorites(st.session_state.username)
    if not lessons:
        st.info("لم تضف أي درس إلى المفضلة.")
    for lesson in lessons:
        _render_lesson_card(lesson)


# =========================================================
# فضاء المطور
# =========================================================

def render_developer_panel():
    st.title("🛠️ فضاء المطور")
    st.caption("إضافة الدروس والأسئلة وإدارة محتوى المنصة.")

    tab1, tab2, tab3 = st.tabs([
        "➕ إضافة درس", "📝 إدارة الأسئلة", "📚 إدارة الدروس"
    ])

    with tab1:
        with st.form("add_lesson_form"):
            subject = st.selectbox(
                "المادة", list(SUBJECTS.keys()),
                format_func=subject_label
            )
            language = st.selectbox(
                "لغة الدرس", ["ar", "fr", "en"],
                format_func=lambda x: {
                    "ar": "العربية", "fr": "الفرنسية", "en": "الإنجليزية"
                }[x]
            )
            title = st.text_input("عنوان الدرس")
            content = st.text_area("محتوى الدرس", height=250)
            image_url = st.text_input("رابط صورة (اختياري)")
            uploaded = st.file_uploader(
                "📎 تحميل ملف الدرس",
                type=["pdf", "txt", "docx", "png", "jpg", "jpeg"]
            )
            auto_extract = st.checkbox(
                "استخراج النص من الملف وإضافته إلى المحتوى",
                value=True
            )
            submit = st.form_submit_button("حفظ الدرس")

        if submit:
            if not title.strip():
                st.error("عنوان الدرس مطلوب.")
            else:
                file_path = ""
                if uploaded is not None:
                    file_path = upload_file(uploaded)

                    if auto_extract and pathlib.Path(file_path).suffix.lower() in (
                        ".pdf", ".txt", ".docx"
                    ):
                        extracted = extract_text_from_file(file_path)
                        if extracted.strip():
                            content = (
                                content.strip() + "\n\n" + extracted
                            ).strip()
                        else:
                            st.warning(
                                "لم يتم استخراج نص. قد يكون الملف صورة ممسوحة ضوئياً."
                            )

                    if pathlib.Path(file_path).suffix.lower() in (
                        ".png", ".jpg", ".jpeg"
                    ):
                        image_url = file_path

                lid = add_lesson(
                    subject, language, title.strip(), content,
                    image_url, file_path, st.session_state.username
                )
                st.success(f"تم حفظ الدرس. رقمه: {lid}")
                st.session_state.current_lesson_id = lid
                st.rerun()

    with tab2:
        lessons = load_lessons()
        if not lessons:
            st.info("أضف درساً أولاً.")
        else:
            lid = st.selectbox(
                "اختر الدرس",
                [l["id"] for l in lessons],
                format_func=lambda x: next(
                    (l["title"] for l in lessons if l["id"] == x), str(x)
                ),
                key="developer_question_lesson"
            )
            st.subheader("➕ إضافة سؤال يدوياً")
            with st.form("add_question_form"):
                question = st.text_area("السؤال")
                a = st.text_input("الخيار A")
                b = st.text_input("الخيار B")
                c = st.text_input("الخيار C")
                d = st.text_input("الخيار D")
                correct = st.selectbox("الجواب الصحيح", ["A", "B", "C", "D"])
                explanation = st.text_area("شرح الإجابة")
                submit_q = st.form_submit_button("إضافة السؤال")

            if submit_q:
                if all(x.strip() for x in [question, a, b, c, d]):
                    add_question(
                        lid, question, a, b, c, d, correct, explanation
                    )
                    st.success("تمت إضافة السؤال.")
                    st.rerun()
                else:
                    st.error("أكمل السؤال وجميع الخيارات.")

            st.divider()
            st.subheader("🤖 اقتراح أسئلة من محتوى الدرس")
            lesson = get_lesson_by_id(lid)
            st.caption(
                "يقترح التطبيق أسئلة محلية اعتماداً على النص. "
                "راجع كل سؤال وخياراته قبل حفظه."
            )
            count = st.slider(
                "عدد الأسئلة المقترحة", 1, 30, 10,
                key="generated_count"
            )
            if st.button("✨ توليد أسئلة", key="generate_questions"):
                generated = generate_questions_from_content(
                    lesson.get("content", ""), count
                )
                st.session_state.generated_questions = generated
                st.session_state.generated_lesson_id = lid

            if (
                st.session_state.get("generated_questions") is not None
                and st.session_state.get("generated_lesson_id") == lid
            ):
                generated = st.session_state.generated_questions
                if not generated:
                    st.warning(
                        "لم يتم العثور على نص كافٍ. أضف محتوى نصياً واضحاً."
                    )
                else:
                    st.write(f"عدد الأسئلة المقترحة: {len(generated)}")
                    with st.form("review_generated_questions"):
                        approved = []
                        for i, q in enumerate(generated):
                            st.markdown(f"### سؤال {i+1}")
                            q_text = st.text_input(
                                "نص السؤال",
                                q["question"],
                                key=f"gen_q_{lid}_{i}"
                            )
                            options = q["options"]
                            a_text = st.text_input(
                                "A", options[0], key=f"gen_a_{lid}_{i}"
                            )
                            b_text = st.text_input(
                                "B", options[1], key=f"gen_b_{lid}_{i}"
                            )
                            c_text = st.text_input(
                                "C", options[2], key=f"gen_c_{lid}_{i}"
                            )
                            d_text = st.text_input(
                                "D", options[3], key=f"gen_d_{lid}_{i}"
                            )
                            correct = st.selectbox(
                                "الإجابة الصحيحة",
                                ["A", "B", "C", "D"],
                                index=["A", "B", "C", "D"].index(q["correct"]),
                                key=f"gen_correct_{lid}_{i}"
                            )
                            explanation = st.text_area(
                                "التفسير",
                                q["explanation"],
                                key=f"gen_explanation_{lid}_{i}"
                            )
                            approved.append({
                                "question": q_text,
                                "a": a_text,
                                "b": b_text,
                                "c": c_text,
                                "d": d_text,
                                "correct": correct,
                                "explanation": explanation
                            })
                        save = st.form_submit_button(
                            "💾 اعتماد الأسئلة وحفظها"
                        )

                    if save:
                        saved = 0
                        for q in approved:
                            if all(q[k].strip() for k in [
                                "question", "a", "b", "c", "d"
                            ]):
                                add_question(
                                    lid, q["question"], q["a"], q["b"],
                                    q["c"], q["d"], q["correct"],
                                    q["explanation"]
                                )
                                saved += 1
                        st.session_state.generated_questions = None
                        st.success(f"تم حفظ {saved} سؤال.")
                        st.rerun()

            st.divider()
            st.subheader("🗑️ الأسئلة الحالية")
            questions = load_questions(lid)
            for q in questions:
                with st.container(border=True):
                    st.write(q["question"])
                    st.caption(
                        f"A: {q['option_a']} | B: {q['option_b']} | "
                        f"C: {q['option_c']} | D: {q['option_d']}"
                    )
                    st.write(f"الإجابة: {q['correct_answer']}")
                    if st.button("حذف السؤال", key=f"delq_{q['id']}"):
                        delete_question(q["id"])
                        st.rerun()

    with tab3:
        lessons = load_lessons()
        for lesson in lessons:
            with st.container(border=True):
                st.write(
                    f"**{lesson['title']}** — "
                    f"{subject_label(lesson['subject'])}"
                )
                if st.button("حذف الدرس", key=f"deletelesson_{lesson['id']}"):
                    if lesson["owner"] == st.session_state.username or (
                        st.session_state.role == "developer"
                    ):
                        delete_lesson(lesson["id"])
                        st.success("تم حذف الدرس.")
                        st.rerun()
                    else:
                        st.error("ليس لديك صلاحية حذف هذا الدرس.")


# =========================================================
# التوجيه الجانبي
# =========================================================

def render_sidebar():
    with st.sidebar:
        st.title("📚 RevisioMaroc")
        st.caption(
            f"مرحباً {st.session_state.get('full_name', '')}"
        )

        theme_name = st.selectbox(
            "🎨 الثيم", list(THEMES.keys()),
            index=list(THEMES.keys()).index(
                st.session_state.get("theme", "🌙 Moonlight")
            ),
            key="sidebar_theme"
        )
        st.session_state.theme = theme_name
        apply_theme()

        language = st.selectbox(
            "🌐 اللغة",
            ["ar", "fr", "en"],
            format_func=lambda x: {
                "ar": "العربية", "fr": "Français", "en": "English"
            }[x],
            index=["ar", "fr", "en"].index(
                st.session_state.get("language", "ar")
            )
        )
        st.session_state.language = language

        st.divider()
        pages = [
            ("dashboard", "🏠 الرئيسية"),
            ("lessons", "📚 الدروس"),
            ("quiz", "📝 الاختبارات"),
            ("review", "⚡ المراجعة السريعة"),
            ("flashcards", "🗂️ بطاقات الحفظ"),
            ("plan", "🗓️ خطة الدراسة"),
            ("friends", "👥 الزملاء"),
            ("leaderboard", "🏆 المتصدرون"),
            ("weekly", "📊 التقرير الأسبوعي"),
            ("stats", "📈 إحصائياتي"),
            ("favorites", "⭐ المفضلة"),
            ("notifications", "🔔 الإشعارات"),
        ]

        if st.session_state.role == "developer":
            pages.append(("developer", "🛠️ فضاء المطور"))

        for key, label in pages:
            if st.button(
                label, key=f"nav_{key}", use_container_width=True
            ):
                st.session_state.page = key
                st.rerun()

        st.divider()
        st.caption(f"المستخدم: {st.session_state.username}")
        if st.button("🚪 تسجيل الخروج", use_container_width=True):
            for key in list(st.session_state.keys()):
                del st.session_state[key]
            st.session_state.authenticated = False
            st.session_state.auth_page = "home"
            st.rerun()


def render_router():
    page = st.session_state.get("page", "dashboard")

    routes = {
        "dashboard": render_dashboard,
        "lessons": render_lessons,
        "lesson_detail": render_lesson_detail,
        "quiz": render_quiz,
        "review": render_quick_review,
        "flashcards": render_flashcards,
        "plan": render_study_plan,
        "friends": render_friends,
        "leaderboard": render_leaderboard,
        "weekly": render_weekly_report,
        "stats": render_my_stats,
        "favorites": render_favorites,
        "notifications": render_notifications,
        "developer": render_developer_panel,
    }

    if page == "developer" and st.session_state.role != "developer":
        st.error("هذه الصفحة متاحة للمطور فقط.")
        st.session_state.page = "dashboard"
        st.rerun()

    routes.get(page, render_dashboard)()
    render_footer()


# =========================================================
# نقطة تشغيل التطبيق
# =========================================================

init_db()

if "authenticated" not in st.session_state:
    st.session_state.authenticated = False

if "auth_page" not in st.session_state:
    st.session_state.auth_page = "home"

if "theme" not in st.session_state:
    st.session_state.theme = "🌙 Moonlight"

if "language" not in st.session_state:
    st.session_state.language = "ar"

if "page" not in st.session_state:
    st.session_state.page = "dashboard"

if "selected_subject" not in st.session_state:
    st.session_state.selected_subject = None

if not st.session_state.authenticated:
    apply_theme()
    render_auth_page()
    st.stop()

apply_theme()
render_sidebar()
render_router()
